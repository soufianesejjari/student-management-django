from rest_framework import views, permissions
from rest_framework.response import Response
from users.models import StudentProfile, TeacherProfile
from academics.models import AcademicYear, Course, Enrollment
from finances.models import Payment, Expense
from planning.models import ClassSession
from django.db.models import Sum, Count, Q
from django.utils import timezone
from datetime import date, datetime, timedelta
from decimal import Decimal
from django.http import HttpResponse


def decimal_to_float(value):
    return float(value or Decimal('0'))


def month_starts(start_date, end_date):
    current = start_date.replace(day=1)
    while current <= end_date:
        yield current
        if current.month == 12:
            current = current.replace(year=current.year + 1, month=1)
        else:
            current = current.replace(month=current.month + 1)

class DashboardStatsView(views.APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        today = date.today()
        current_month = today.month
        current_year = today.year
        academic_year = AcademicYear.get_active()

        # Active Students
        active_students = StudentProfile.objects.filter(status='ACTIVE').count()
        
        # Active Teachers
        active_teachers = TeacherProfile.objects.filter(status='ACTIVE').count()
        
        # Active Courses
        active_courses = Course.objects.filter(status='ACTIVE').count()

        # Monthly Revenue
        monthly_income = Payment.objects.filter(
            date__month=current_month, 
            date__year=current_year,
            status='PAID'
        ).filter(
            Q(subscription__enrollment__academic_year=academic_year) |
            Q(subscription__isnull=True)
        ).aggregate(Sum('amount'))['amount__sum'] or 0

        # Monthly Expenses
        monthly_expenses = Expense.objects.filter(
            date__month=current_month, 
            date__year=current_year
        ).aggregate(Sum('amount'))['amount__sum'] or 0

        # Calculations
        net_revenue = monthly_income - monthly_expenses

        # Recent Enrollments (This month)
        new_students = StudentProfile.objects.filter(
            enrollment_date__month=current_month,
            enrollment_date__year=current_year
        ).count()

        return Response({
            'active_students': active_students,
            'active_teachers': active_teachers,
            'active_courses': active_courses,
            'monthly_revenue': monthly_income,
            'monthly_expenses': monthly_expenses,
            'net_revenue': net_revenue,
            'new_students_this_month': new_students
        })

class UpcomingClassesView(views.APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        """Return the next actual occurrences from recurring sessions."""
        academic_year = AcademicYear.get_active()
        now = timezone.localtime()
        today = now.date()
        window_end = min(today + timedelta(days=30), academic_year.end_date)
        sessions = ClassSession.objects.filter(
            academic_year=academic_year,
            start_date__lte=window_end,
        ).filter(
            Q(end_date__gte=today) | Q(end_date__isnull=True)
        ).select_related('course', 'teacher__user', 'room')

        occurrences = []
        for session in sessions:
            occurrence_date = today + timedelta(days=(session.day_of_week - today.weekday()) % 7)
            if occurrence_date < session.start_date:
                days_until_session_day = (session.day_of_week - session.start_date.weekday()) % 7
                occurrence_date = session.start_date + timedelta(days=days_until_session_day)

            effective_end = session.end_date or academic_year.end_date
            if occurrence_date > effective_end or occurrence_date > window_end:
                continue

            occurrence_start = timezone.make_aware(datetime.combine(occurrence_date, session.start_time))
            if occurrence_start <= now:
                occurrence_date += timedelta(days=7)
                occurrence_start = timezone.make_aware(datetime.combine(occurrence_date, session.start_time))

            if occurrence_date > effective_end or occurrence_date > window_end:
                continue

            occurrences.append({
                'id': session.id,
                'course': session.course.name,
                'teacher': session.teacher.user.get_full_name() or session.teacher.user.username,
                'room': session.room.name,
                'date': occurrence_date.isoformat(),
                'day': session.day_of_week,
                'start': session.start_time.strftime('%H:%M'),
                'end': session.end_time.strftime('%H:%M'),
                '_sort': occurrence_start,
            })

        occurrences.sort(key=lambda item: item['_sort'])
        for item in occurrences:
            item.pop('_sort', None)

        return Response(occurrences[:5])


class ReportsView(views.APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        academic_year = AcademicYear.get_active()
        start = academic_year.start_date
        end = academic_year.end_date

        enrollments = Enrollment.objects.filter(academic_year=academic_year)
        active_enrollments = enrollments.filter(status='ACTIVE')
        paid_payments = Payment.objects.filter(status='PAID', date__range=(start, end)).filter(
            Q(subscription__enrollment__academic_year=academic_year) |
            Q(subscription__isnull=True)
        ).select_related('subscription__enrollment__course__subject')
        paid_expenses = Expense.objects.filter(status='PAID', date__range=(start, end))

        monthly_rows = []
        for month_start in month_starts(start, end):
            if month_start.month == 12:
                next_month = month_start.replace(year=month_start.year + 1, month=1)
            else:
                next_month = month_start.replace(month=month_start.month + 1)
            month_end = min(next_month - timedelta(days=1), end)

            monthly_rows.append({
                'month': month_start.strftime('%Y-%m'),
                'label': month_start.strftime('%b %Y'),
                'enrollments': enrollments.filter(enrolled_at__date__range=(month_start, month_end)).count(),
                'revenue': decimal_to_float(paid_payments.filter(date__range=(month_start, month_end)).aggregate(total=Sum('amount'))['total']),
                'expenses': decimal_to_float(paid_expenses.filter(date__range=(month_start, month_end)).aggregate(total=Sum('amount'))['total']),
            })

        revenue_by_subject = {}
        for payment in paid_payments:
            subject = 'Autre'
            if payment.subscription_id:
                enrollment = payment.subscription.enrollment
                subject = enrollment.course.subject.name if enrollment and enrollment.course.subject_id else 'Autre'
            revenue_by_subject[subject] = revenue_by_subject.get(subject, Decimal('0')) + payment.amount

        expense_by_category = {
            row['category']: decimal_to_float(row['total'])
            for row in paid_expenses.values('category').annotate(total=Sum('amount')).order_by('-total')
        }

        def ordered_items(mapping):
            return [
                {'label': label, 'value': decimal_to_float(value)}
                for label, value in sorted(mapping.items(), key=lambda item: item[1], reverse=True)
            ]

        age_distribution = [
            {'label': row['age_group'] or 'Non renseigné', 'value': row['total']}
            for row in StudentProfile.objects.values('age_group').annotate(total=Count('id')).order_by('-total')
        ]

        instrument_distribution = [
            {'label': row['course__subject__name'] or 'Non renseigné', 'value': row['total']}
            for row in active_enrollments.values('course__subject__name').annotate(total=Count('id')).order_by('-total')
        ]

        course_popularity = [
            {'label': row['course__name'] or 'Non renseigné', 'value': row['total']}
            for row in active_enrollments.values('course__name').annotate(total=Count('id')).order_by('-total')[:10]
        ]

        retention = [
            {'label': row['status'], 'value': row['total']}
            for row in enrollments.values('status').annotate(total=Count('id')).order_by('-total')
        ]

        room_usage = [
            {'label': row['room__name'] or 'Non renseigné', 'value': row['total']}
            for row in ClassSession.objects.filter(academic_year=academic_year)
            .values('room__name').annotate(total=Count('id')).order_by('-total')
        ]

        total_revenue = decimal_to_float(paid_payments.aggregate(total=Sum('amount'))['total'])
        total_expenses = decimal_to_float(paid_expenses.aggregate(total=Sum('amount'))['total'])

        return Response({
            'academic_year': {
                'id': academic_year.id,
                'name': academic_year.name,
                'start_date': start.isoformat(),
                'end_date': end.isoformat(),
            },
            'summary': {
                'students': StudentProfile.objects.filter(status='ACTIVE').count(),
                'teachers': TeacherProfile.objects.filter(status='ACTIVE').count(),
                'courses': Course.objects.filter(status='ACTIVE').count(),
                'active_enrollments': active_enrollments.count(),
                'total_revenue': total_revenue,
                'total_expenses': total_expenses,
                'net_revenue': total_revenue - total_expenses,
            },
            'monthly': monthly_rows,
            'revenue_by_subject': ordered_items(revenue_by_subject),
            'expense_by_category': [
                {'label': label, 'value': value}
                for label, value in expense_by_category.items()
            ],
            'age_distribution': age_distribution,
            'instrument_distribution': instrument_distribution,
            'course_popularity': course_popularity,
            'retention': retention,
            'room_usage': room_usage,
        })


class ReportsExcelExportView(views.APIView):
    """Export the academy's finance report in an editable Excel workbook."""
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from planning.models import TeacherMonthlyPayroll

        academic_year = AcademicYear.get_active()
        start, end = academic_year.start_date, academic_year.end_date
        paid_payments = Payment.objects.filter(status='PAID', date__range=(start, end)).select_related('student__user')
        payrolls = TeacherMonthlyPayroll.objects.filter(academic_year=academic_year).select_related('teacher__user', 'expense')

        workbook = Workbook()
        summary = workbook.active
        summary.title = 'Synthèse'
        student_sheet = workbook.create_sheet('Paiements élèves')
        salary_sheet = workbook.create_sheet('Salaires professeurs')

        header_fill = PatternFill('solid', fgColor='D8262B')
        header_font = Font(color='FFFFFF', bold=True)

        def write_header(sheet, columns):
            sheet.append(columns)
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center')
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions

        income = paid_payments.aggregate(total=Sum('amount'))['total'] or Decimal('0')
        expenses = Expense.objects.filter(status='PAID', date__range=(start, end)).aggregate(total=Sum('amount'))['total'] or Decimal('0')
        write_header(summary, ['Année académique', 'Date début', 'Date fin', 'Encaissements (MAD)', 'Dépenses (MAD)', 'Résultat net (MAD)'])
        summary.append([academic_year.name, start, end, float(income), float(expenses), float(income - expenses)])

        write_header(student_sheet, ['Élève', 'Mois', 'Nombre de paiements', 'Montant payé (MAD)'])
        student_rows = paid_payments.values(
            'student__user__first_name', 'student__user__last_name', 'student__user__username', 'date__year', 'date__month'
        ).annotate(payment_count=Count('id'), amount=Sum('amount')).order_by(
            'student__user__last_name', 'student__user__first_name', 'date__year', 'date__month'
        )
        for row in student_rows:
            name = f"{row['student__user__first_name']} {row['student__user__last_name']}".strip() or row['student__user__username']
            student_sheet.append([name, f"{row['date__year']}-{row['date__month']:02d}", row['payment_count'], float(row['amount'])])

        write_header(salary_sheet, ['Professeur', 'Mois', 'Statut', 'Heures travaillées', 'Taux horaire (MAD)', 'Salaire (MAD)', 'Dépense payée'])
        for payroll in payrolls:
            teacher = payroll.teacher.user.get_full_name().strip() or payroll.teacher.user.username
            salary_sheet.append([
                teacher, f'{payroll.year}-{payroll.month:02d}', payroll.status, float(payroll.worked_hours),
                float(payroll.hourly_rate), float(payroll.amount), payroll.expense.status if payroll.expense_id else 'NON CRÉÉE',
            ])

        for sheet in workbook.worksheets:
            for column_cells in sheet.columns:
                sheet.column_dimensions[column_cells[0].column_letter].width = min(max(len(str(cell.value or '')) for cell in column_cells) + 2, 32)
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = '#,##0.00'

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="rapport-financier-{academic_year.name}.xlsx"'
        workbook.save(response)
        return response
